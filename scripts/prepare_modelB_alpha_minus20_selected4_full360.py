"""Validate four real alpha=-20 full360 datasets and build a one-sheet XLSX."""

import argparse
import csv
import math
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
Z_ORDER = [111.0, 94.0, 88.0, 82.0]
PHI = [round(i * 3.6, 10) for i in range(101)]
CLOSURE_FIELDS = [
    "Fx_total_corr_mN",
    "DeltaFx_ball_mN",
    "Fy_total_corr_mN",
    "Fz_total_corr_mN",
]
XLSX_COLUMNS = [
    "z_sphere_mm",
    "alpha_deg",
    "phi_deg",
    "Fx_total_corr_mN",
    "Fy_total_corr_mN",
    "Fz_total_corr_mN",
    "Fx_hold_corr_mN",
    "DeltaFx_ball_mN",
    "Fx_B0_raw_mN",
    "Fx_B1_raw_mN",
    "Fx_B2_raw_mN",
    "elements",
    "DOF",
    "min_quality",
    "same_mesh_verified",
    "status",
]
NUMERIC_COLUMNS = set(XLSX_COLUMNS) - {"same_mesh_verified", "status"}
INTEGER_COLUMNS = {"elements", "DOF"}


def fail(message):
    raise ValueError(message)


def number(row, column):
    try:
        value = float(row[column])
    except (KeyError, TypeError, ValueError):
        fail("invalid numeric field {}={!r}".format(column, row.get(column)))
    if not math.isfinite(value):
        fail("non-finite numeric field {}={!r}".format(column, row.get(column)))
    return value


def read_group(path, z):
    if not path.is_file():
        fail("required source CSV is missing: {}".format(path))
    with path.open("r", newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames:
            fail("CSV has no header: {}".format(path))
        rows = [row for row in reader if abs(number(row, "z_sphere_mm") - z) < 1e-8]
    if len(rows) != 101:
        fail("z={} in {} has {} rows, expected 101".format(z, path.name, len(rows)))
    validate_group(rows, z, path.name)
    return rows


def validate_group(rows, z, label):
    if len(rows) != 101:
        fail("{} z={} has {} rows, expected 101".format(label, z, len(rows)))
    phis = [number(row, "phi_deg") for row in rows]
    if len(set(phis)) != 101 or any(abs(a - b) > 1e-8 for a, b in zip(sorted(phis), PHI)):
        fail("{} z={} does not contain unique phi=0:3.6:360".format(label, z))
    rows.sort(key=lambda row: number(row, "phi_deg"))
    for row in rows:
        if abs(number(row, "z_sphere_mm") - z) > 1e-8:
            fail("{} has a mixed height".format(label))
        if abs(number(row, "alpha_deg") + 20.0) > 1e-10:
            fail("{} has alpha other than -20 deg".format(label))
        if row.get("status") != "SUCCESS":
            fail("{} has non-SUCCESS row at phi={}".format(label, row.get("phi_deg")))
        if str(row.get("same_mesh_verified", "")).strip().lower() != "true":
            fail("{} has unverified same-mesh row at phi={}".format(label, row.get("phi_deg")))
        for column in XLSX_COLUMNS:
            if column not in row:
                fail("{} is missing required column {}".format(label, column))
        fx0 = number(row, "Fx_B0_raw_mN")
        fx1 = number(row, "Fx_B1_raw_mN")
        fx2 = number(row, "Fx_B2_raw_mN")
        if abs(number(row, "Fx_hold_corr_mN") - (fx1 - fx0)) > 2e-8:
            fail("{} Fx_hold_corr mismatch at phi={}".format(label, row["phi_deg"]))
        if abs(number(row, "DeltaFx_ball_mN") - (fx2 - fx1)) > 2e-8:
            fail("{} DeltaFx_ball mismatch at phi={}".format(label, row["phi_deg"]))
        if abs(number(row, "Fx_total_corr_mN") - (fx2 - fx0)) > 2e-8:
            fail("{} Fx_total_corr mismatch at phi={}".format(label, row["phi_deg"]))


def closure_error(rows):
    at0 = next(row for row in rows if abs(number(row, "phi_deg")) < 1e-8)
    at360 = next(row for row in rows if abs(number(row, "phi_deg") - 360.0) < 1e-8)
    return max(abs(number(at0, key) - number(at360, key)) for key in CLOSURE_FIELDS)


def load_existing_groups():
    critical = DATA / "modelB_alpha_minus20_critical_full360.csv"
    levels = DATA / "modelB_alpha_minus20_force_levels_full360.csv"
    return {
        111.0: read_group(critical, 111.0),
        94.0: read_group(levels, 94.0),
        82.0: read_group(levels, 82.0),
    }


def validate_prescan_88(rows):
    matches = read_prescan_88()
    full90 = next(row for row in rows if abs(number(row, "phi_deg") - 90.0) < 1e-8)
    delta = abs(number(full90, "Fx_total_corr_mN") - number(matches[0], "Fx_total_corr_mN"))
    if delta > 1e-5:
        fail("z=88 phi=90 differs from prescan by {:.12g} mN".format(delta))
    return number(matches[0], "Fx_total_corr_mN"), number(full90, "Fx_total_corr_mN"), delta


def read_prescan_88():
    path = DATA / "modelB_alpha_minus20_force_level_prescan.csv"
    if not path.is_file():
        fail("z=88 phi=90 prescan CSV is missing: {}".format(path))
    with path.open("r", newline="", encoding="utf-8-sig") as stream:
        prescan = csv.DictReader(stream)
        matches = [
            row for row in prescan
            if abs(number(row, "z_sphere_mm") - 88.0) < 1e-8
            and abs(number(row, "phi_deg") - 90.0) < 1e-8
        ]
    if len(matches) != 1:
        fail("expected exactly one existing z=88 phi=90 prescan row, found {}".format(len(matches)))
    if abs(number(matches[0], "alpha_deg") + 20.0) > 1e-10:
        fail("z=88 phi=90 prescan alpha is not -20 deg")
    if matches[0].get("status") != "SUCCESS" or str(matches[0].get("same_mesh_verified", "")).lower() != "true":
        fail("z=88 phi=90 prescan point is not validated same-mesh SUCCESS")
    expected = 0.146497737406
    if abs(number(matches[0], "Fx_total_corr_mN") - expected) > 1e-8:
        fail("existing z=88 phi=90 prescan differs from expected sanity value")
    return matches


def values_for_excel(rows):
    result = []
    for row in rows:
        values = []
        for column in XLSX_COLUMNS:
            if column == "same_mesh_verified":
                values.append(True)
            elif column in INTEGER_COLUMNS:
                values.append(int(round(number(row, column))))
            elif column in NUMERIC_COLUMNS:
                values.append(number(row, column))
            else:
                values.append(row[column])
        result.append(values)
    return result


def validate_workbook(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != ["full360"]:
            fail("Excel sheets must be exactly ['full360'], got {}".format(book.sheetnames))
        sheet = book["full360"]
        if sheet.max_row != 405 or sheet.max_column != len(XLSX_COLUMNS):
            fail("Excel dimensions are {}x{}, expected 405x{}".format(sheet.max_row, sheet.max_column, len(XLSX_COLUMNS)))
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if header != XLSX_COLUMNS:
            fail("Excel header does not match required long-format columns")
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        if len(rows) != 404:
            fail("Excel has {} data rows, expected 404".format(len(rows)))
        z_col = XLSX_COLUMNS.index("z_sphere_mm")
        alpha_col = XLSX_COLUMNS.index("alpha_deg")
        phi_col = XLSX_COLUMNS.index("phi_deg")
        mesh_col = XLSX_COLUMNS.index("same_mesh_verified")
        status_col = XLSX_COLUMNS.index("status")
        actual_z = [float(row[z_col]) for row in rows]
        expected_z = [z for z in Z_ORDER for _ in PHI]
        if actual_z != expected_z:
            fail("Excel z order/count differs from [111,94,88,82], each 101 rows")
        for index, z in enumerate(Z_ORDER):
            group = rows[index * 101:(index + 1) * 101]
            if any(abs(float(row[alpha_col]) + 20.0) > 1e-10 for row in group):
                fail("Excel has alpha other than -20 at z={}".format(z))
            if any(abs(float(row[phi_col]) - expected) > 1e-8 for row, expected in zip(group, PHI)):
                fail("Excel phase sequence invalid at z={}".format(z))
            if any(row[mesh_col] is not True for row in group):
                fail("Excel same_mesh_verified is not TRUE at z={}".format(z))
            if any(row[status_col] != "SUCCESS" for row in group):
                fail("Excel contains non-SUCCESS data at z={}".format(z))
    finally:
        book.close()


def print_force_summaries(groups):
    for z in Z_ORDER:
        rows = groups[z]
        fmax = max(rows, key=lambda row: number(row, "Fx_total_corr_mN"))
        fmin = min(rows, key=lambda row: number(row, "Fx_total_corr_mN"))
        print("z={:g} mm Fmax={:.12g} mN @ {:.12g} deg; Fmin={:.12g} mN @ {:.12g} deg".format(
            z,
            number(fmax, "Fx_total_corr_mN"), number(fmax, "phi_deg"),
            number(fmin, "Fx_total_corr_mN"), number(fmin, "phi_deg"),
        ))


def main():
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--preflight", action="store_true")
    mode.add_argument("--finalize", action="store_true")
    parser.add_argument("--run-csv", type=Path)
    parser.add_argument("--candidate-xlsx", type=Path)
    args = parser.parse_args()

    z88_csv = DATA / "modelB_alpha_minus20_z88_full360.csv"
    final_xlsx = DATA / "modelB_alpha_minus20_selected4_full360.xlsx"
    if args.preflight:
        if z88_csv.exists() or final_xlsx.exists():
            fail("final z88 CSV or selected-four XLSX already exists; refusing overwrite")
        existing = load_existing_groups()
        prescan = read_prescan_88()
        print("PRECHECK_OK z111=101 z94=101 z82=101; z88_targets_absent=true")
        for z in (111.0, 94.0, 82.0):
            print("EXISTING z={:g} closure_max={:.12g} mN".format(z, closure_error(existing[z])))
        print("PRESCAN z=88 phi=90 Fx_total_corr={:.12g} mN".format(number(prescan[0], "Fx_total_corr_mN")))
        return

    if args.run_csv is None or args.candidate_xlsx is None:
        fail("--finalize requires --run-csv and --candidate-xlsx")
    if z88_csv.exists() or final_xlsx.exists():
        fail("final z88 CSV or selected-four XLSX already exists; refusing overwrite")
    existing = load_existing_groups()
    z88 = read_group(args.run_csv, 88.0)
    prescan90, full90, delta90 = validate_prescan_88(z88)
    baseline = max(closure_error(existing[z]) for z in (111.0, 94.0, 82.0))
    z88_closure = closure_error(z88)
    closure_limit = max(10.0 * baseline, 1e-8)
    if z88_closure > closure_limit:
        fail("z=88 max closure {:.12g} mN exceeds {:.12g} mN baseline limit".format(z88_closure, closure_limit))
    print("Z88_ACCEPTED rows=101 alpha=-20 closure_max={:.12g} mN limit={:.12g} mN".format(z88_closure, closure_limit))
    print("Z88_PHI90 prescan={:.12g} full360={:.12g} abs_delta={:.12g} mN".format(prescan90, full90, delta90))

    groups = dict(existing)
    groups[88.0] = z88
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "full360"
    sheet.append(XLSX_COLUMNS)
    for z in Z_ORDER:
        for values in values_for_excel(groups[z]):
            sheet.append(values)
    args.candidate_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.candidate_xlsx)
    workbook.close()
    validate_workbook(args.candidate_xlsx)

    shutil.copyfile(args.run_csv, z88_csv)
    shutil.copyfile(args.candidate_xlsx, final_xlsx)
    validate_workbook(final_xlsx)
    print_force_summaries(groups)
    print("SELECTED4_XLSX_OK rows=404 groups=111,94,88,82 sheets=full360")
    print("Z88_CSV={}".format(z88_csv))
    print("XLSX={}".format(final_xlsx))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print("VALIDATION_ERROR: {}".format(error), file=sys.stderr)
        raise
