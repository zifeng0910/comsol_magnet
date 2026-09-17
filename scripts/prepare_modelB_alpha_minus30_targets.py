"""Select measured alpha=-30 heights for four alpha=-20 force targets and validate full360 data."""
from __future__ import annotations

import argparse
import csv
import itertools
import math
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
TARGET_BOOK = DATA / "modelB_alpha_minus20_selected4_full360.xlsx"
TARGET_Z_ORDER = [111.0, 94.0, 88.0, 82.0]
COARSE_ZS = [float(z) for z in range(72, 125, 4)]
PHI = [round(i * 3.6, 10) for i in range(101)]
FORCE_FIELDS = ["Fx_total_corr_mN", "DeltaFx_ball_mN", "Fy_total_corr_mN", "Fz_total_corr_mN"]
RAW_FIELDS = ["Fx_B0_raw_mN", "Fx_B1_raw_mN", "Fx_B2_raw_mN"]
PRESCAN_COLUMNS = [
    "z_sphere_mm", "alpha_deg", "phi_deg", "Fx_B0_raw_mN", "Fx_B1_raw_mN", "Fx_B2_raw_mN",
    "Fx_hold_corr_mN", "DeltaFx_ball_mN", "Fx_total_corr_mN", "Fy_total_corr_mN", "Fz_total_corr_mN",
    "elements", "DOF", "min_quality", "same_mesh_verified", "status",
]
FULL_COLUMNS = PRESCAN_COLUMNS
EXCEL_COLUMNS = [
    "z_sphere_mm", "alpha_deg", "phi_deg", "Fx_total_corr_mN", "Fy_total_corr_mN", "Fz_total_corr_mN",
    "Fx_hold_corr_mN", "DeltaFx_ball_mN", "Fx_B0_raw_mN", "Fx_B1_raw_mN", "Fx_B2_raw_mN",
    "elements", "DOF", "min_quality", "same_mesh_verified", "status",
]
NUMERIC_EXCEL = set(EXCEL_COLUMNS) - {"same_mesh_verified", "status"}
INTEGER_EXCEL = {"elements", "DOF"}


def fail(message: str) -> None:
    raise ValueError(message)


def num(row: dict, key: str) -> float:
    try:
        value = float(row[key])
    except (KeyError, TypeError, ValueError):
        fail("invalid {}={!r}".format(key, row.get(key)))
    if not math.isfinite(value):
        fail("non-finite {}={!r}".format(key, row.get(key)))
    return value


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        fail("required CSV missing: {}".format(path))
    with path.open("r", newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames:
            fail("CSV has no header: {}".format(path))
        return list(reader)


def target_levels() -> list[dict[str, float]]:
    if not TARGET_BOOK.is_file():
        fail("previously accepted alpha=-20 target workbook missing")
    book = load_workbook(TARGET_BOOK, read_only=True, data_only=True)
    try:
        if book.sheetnames != ["full360"]:
            fail("alpha=-20 reference workbook must only contain full360")
        sheet = book["full360"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    header = list(rows[0])
    if "z_sphere_mm" not in header or "Fx_total_corr_mN" not in header:
        fail("alpha=-20 reference workbook columns are invalid")
    zi, fi = header.index("z_sphere_mm"), header.index("Fx_total_corr_mN")
    grouped: dict[float, list[float]] = {}
    for row in rows[1:]:
        grouped.setdefault(float(row[zi]), []).append(float(row[fi]))
    if [z for z in TARGET_Z_ORDER if z in grouped] != TARGET_Z_ORDER:
        fail("alpha=-20 reference workbook does not contain selected four heights")
    return [{"target_force_mN": max(grouped[z]), "reference_z_mm": z} for z in TARGET_Z_ORDER]


def validate_prescan_rows(rows: list[dict[str, str]], label: str) -> list[dict[str, str]]:
    seen: set[float] = set()
    for row in rows:
        for key in PRESCAN_COLUMNS:
            if key not in row:
                fail("{} missing column {}".format(label, key))
        z = num(row, "z_sphere_mm")
        if z in seen:
            fail("{} contains duplicate z={}".format(label, z))
        seen.add(z)
        if abs(num(row, "alpha_deg") + 30.0) > 1e-10 or abs(num(row, "phi_deg") - 90.0) > 1e-10:
            fail("{} has wrong alpha/phi at z={}".format(label, z))
        if row["status"] != "SUCCESS" or row["same_mesh_verified"].strip().lower() != "true":
            fail("{} has invalid status/same-mesh at z={}".format(label, z))
        b0, b1, b2 = (num(row, key) for key in RAW_FIELDS)
        for key, expected in (("Fx_hold_corr_mN", b1 - b0), ("DeltaFx_ball_mN", b2 - b1), ("Fx_total_corr_mN", b2 - b0)):
            if abs(num(row, key) - expected) > 2e-8:
                fail("{} force difference mismatch at z={} field={}".format(label, z, key))
    return rows


def combine_prescans(coarse: Path, refine: Path | None = None) -> list[dict[str, str]]:
    rows = validate_prescan_rows(read_csv(coarse), "coarse prescan")
    if refine is not None:
        rows.extend(validate_prescan_rows(read_csv(refine), "refinement prescan"))
    by_z: dict[float, dict[str, str]] = {}
    for row in rows:
        z = num(row, "z_sphere_mm")
        if z in by_z:
            fail("prescan stages overlap at z={}".format(z))
        by_z[z] = row
    return [by_z[z] for z in sorted(by_z)]


def target_brackets(rows: list[dict[str, str]]) -> list[tuple[dict[str, float], dict[str, str], dict[str, str]]]:
    ordered = sorted(rows, key=lambda row: num(row, "z_sphere_mm"))
    result = []
    for target in target_levels():
        value = target["target_force_mN"]
        exact = [r for r in ordered if abs(num(r, "Fx_total_corr_mN") - value) < 1e-12]
        if exact:
            result.append((target, exact[0], exact[0]))
            continue
        pairs = []
        for low_z, high_z in zip(ordered, ordered[1:]):
            low_f, high_f = num(low_z, "Fx_total_corr_mN"), num(high_z, "Fx_total_corr_mN")
            if min(low_f, high_f) <= value <= max(low_f, high_f) and abs(high_f - low_f) > 1e-12:
                z_est = num(low_z, "z_sphere_mm") + (value - low_f) * (num(high_z, "z_sphere_mm") - num(low_z, "z_sphere_mm")) / (high_f - low_f)
                pairs.append((abs(z_est - target["reference_z_mm"]), z_est, low_z, high_z))
        if not pairs:
            observed = [num(r, "Fx_total_corr_mN") for r in ordered]
            fail("target {:.12g} mN is not bracketed; measured range {:.12g}..{:.12g} mN".format(value, min(observed), max(observed)))
        _, _, low_z, high_z = min(pairs, key=lambda item: item[0])
        result.append((target, low_z, high_z))
    return result


def propose_refinement(coarse: Path, output: Path) -> list[dict[str, str]]:
    rows = validate_prescan_rows(read_csv(coarse), "coarse prescan")
    actual = {num(row, "z_sphere_mm") for row in rows}
    if set(COARSE_ZS) - actual:
        fail("coarse scan incomplete; missing z={}".format(sorted(set(COARSE_ZS) - actual)))
    proposals = []
    used: set[float] = set()
    for target, lo, hi in target_brackets(rows):
        f0, f1 = num(lo, "Fx_total_corr_mN"), num(hi, "Fx_total_corr_mN")
        z0, z1 = num(lo, "z_sphere_mm"), num(hi, "z_sphere_mm")
        estimate = z0 if z0 == z1 else z0 + (target["target_force_mN"] - f0) * (z1 - z0) / (f1 - f0)
        candidate = round(estimate * 2.0) / 2.0
        while candidate in used:
            candidate += 0.5
        if candidate < min(z0, z1) or candidate > max(z0, z1):
            candidate = round(estimate * 2.0) / 2.0
            while candidate in used:
                candidate -= 0.5
        used.add(candidate)
        proposals.append({
            "target_force_mN": "{:.12g}".format(target["target_force_mN"]),
            "reference_z_mm": "{:.12g}".format(target["reference_z_mm"]),
            "candidate_z_mm": "{:.12g}".format(candidate),
            "bracket_z1_mm": "{:.12g}".format(z0),
            "bracket_fx1_mN": "{:.12g}".format(f0),
            "bracket_z2_mm": "{:.12g}".format(z1),
            "bracket_fx2_mN": "{:.12g}".format(f1),
            "candidate_already_measured": str(candidate in actual).lower(),
        })
    if len({float(row["candidate_z_mm"]) for row in proposals}) != 4:
        fail("four distinct refinement heights could not be proposed")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(proposals[0]))
        writer.writeheader()
        writer.writerows(proposals)
    for row in proposals:
        print("TARGET {:.9f} mN -> real COMSOL candidate z={} mm (bracket {}..{} mm)".format(
            float(row["target_force_mN"]), row["candidate_z_mm"], row["bracket_z1_mm"], row["bracket_z2_mm"]))
    return proposals


def select_measured_heights(rows: list[dict[str, str]], output: Path) -> list[dict[str, str]]:
    targets = target_levels()
    for target, _, _ in target_brackets(rows):
        del target
    best = None
    for assignment in itertools.permutations(rows, len(targets)):
        errors = [abs(num(row, "Fx_total_corr_mN") - target["target_force_mN"]) for target, row in zip(targets, assignment)]
        score = sum(errors)
        if best is None or score < best[0]:
            best = (score, assignment, errors)
    if best is None:
        fail("not enough distinct measured heights to select four targets")
    selected = []
    for target, row, error in zip(targets, best[1], best[2]):
        selected.append({
            "target_force_mN": "{:.12g}".format(target["target_force_mN"]),
            "reference_z_mm": "{:.12g}".format(target["reference_z_mm"]),
            "selected_z_mm": "{:.12g}".format(num(row, "z_sphere_mm")),
            "measured_phi90_Fx_mN": "{:.12g}".format(num(row, "Fx_total_corr_mN")),
            "selection_abs_error_mN": "{:.12g}".format(error),
        })
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(selected[0]))
        writer.writeheader()
        writer.writerows(selected)
    for row in selected:
        print("SELECTED target={:.9f} mN z={} mm measured_phi90={:.9f} error={:.9f} mN".format(
            float(row["target_force_mN"]), row["selected_z_mm"], float(row["measured_phi90_Fx_mN"]), float(row["selection_abs_error_mN"])))
    return selected


def closure(rows: list[dict[str, str]]) -> dict[str, float]:
    zero = next(row for row in rows if abs(num(row, "phi_deg")) < 1e-8)
    full = next(row for row in rows if abs(num(row, "phi_deg") - 360.0) < 1e-8)
    return {key: abs(num(zero, key) - num(full, key)) for key in FORCE_FIELDS}


def validate_curve(rows: list[dict[str, str]], z: float) -> None:
    if len(rows) != 101:
        fail("z={} has {} rows, expected 101".format(z, len(rows)))
    ordered = sorted(rows, key=lambda row: num(row, "phi_deg"))
    phases = [num(row, "phi_deg") for row in ordered]
    if any(abs(a - b) > 1e-8 for a, b in zip(phases, PHI)):
        fail("z={} phases are not the 101 real points 0:3.6:360".format(z))
    for row in ordered:
        if abs(num(row, "z_sphere_mm") - z) > 1e-8 or abs(num(row, "alpha_deg") + 30.0) > 1e-10:
            fail("z/alpha mismatch in full360 row")
        if row["status"] != "SUCCESS" or row["same_mesh_verified"].strip().lower() != "true":
            fail("unsuccessful or unverified mesh result at z={} phi={}".format(z, row["phi_deg"]))
        b0, b1, b2 = (num(row, key) for key in RAW_FIELDS)
        for key, expected in (("Fx_hold_corr_mN", b1-b0), ("DeltaFx_ball_mN", b2-b1), ("Fx_total_corr_mN", b2-b0)):
            if abs(num(row, key) - expected) > 2e-8:
                fail("force difference mismatch z={} phi={} field={}".format(z, row["phi_deg"], key))


def excel_values(row: dict[str, str]) -> list[object]:
    values = []
    for key in EXCEL_COLUMNS:
        if key == "same_mesh_verified":
            values.append(True)
        elif key in INTEGER_EXCEL:
            values.append(int(round(num(row, key))))
        elif key in NUMERIC_EXCEL:
            values.append(num(row, key))
        else:
            values.append(row[key])
    return values


def validate_workbook(path: Path, z_order: list[float]) -> None:
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != ["full360"]:
            fail("output workbook must contain only full360 sheet")
        sheet = book["full360"]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) != 405 or list(rows[0]) != EXCEL_COLUMNS:
            fail("workbook must have the required header and 404 rows")
        data = rows[1:]
        for group_index, z in enumerate(z_order):
            group = data[group_index*101:(group_index+1)*101]
            if len(group) != 101:
                fail("Excel group size invalid at z={}".format(z))
            zidx, aidx, pidx = EXCEL_COLUMNS.index("z_sphere_mm"), EXCEL_COLUMNS.index("alpha_deg"), EXCEL_COLUMNS.index("phi_deg")
            if any(abs(float(row[zidx])-z)>1e-8 or abs(float(row[aidx])+30)>1e-10 for row in group):
                fail("Excel z/alpha mismatch at group {}".format(z))
            if any(abs(float(row[pidx])-phi)>1e-8 for row,phi in zip(group,PHI)):
                fail("Excel phi sequence mismatch at z={}".format(z))
    finally:
        book.close()


def finalize(args: argparse.Namespace) -> None:
    coarse = validate_prescan_rows(read_csv(args.coarse_csv), "coarse prescan")
    refined = validate_prescan_rows(read_csv(args.refine_csv), "refinement prescan")
    all_pre = combine_prescans(args.coarse_csv, args.refine_csv)
    selected = list(csv.DictReader(args.selected_csv.open("r", newline="", encoding="utf-8-sig")))
    targets = target_levels()
    if len(selected) != 4 or len({float(row["selected_z_mm"]) for row in selected}) != 4:
        fail("selected target table must contain four unique measured heights")
    if [float(row["target_force_mN"]) for row in selected] != [row["target_force_mN"] for row in targets]:
        fail("selected target force levels/order mismatch")

    full = read_csv(args.full_csv)
    summary_source = read_csv(args.run_summary)
    z_order = [float(row["selected_z_mm"]) for row in selected]
    if len(full) != 404 or len(summary_source) != 4:
        fail("full360 rows/summary groups are not 404/4")
    groups: dict[float, list[dict[str, str]]] = {z: [] for z in z_order}
    for row in full:
        z = num(row, "z_sphere_mm")
        if z not in groups:
            fail("unexpected z in full360 data: {}".format(z))
        groups[z].append(row)
    summary_rows = []
    for target_row, target in zip(selected, targets):
        z = float(target_row["selected_z_mm"])
        rows = groups[z]
        validate_curve(rows, z)
        ordered = sorted(rows, key=lambda row: num(row, "phi_deg"))
        closure_values = closure(ordered)
        closure_max = max(closure_values.values())
        at90 = next(row for row in ordered if abs(num(row, "phi_deg")-90.0)<1e-8)
        prescan = next(row for row in all_pre if abs(num(row, "z_sphere_mm")-z)<1e-8)
        delta90 = abs(num(at90, "Fx_total_corr_mN")-num(prescan, "Fx_total_corr_mN"))
        if delta90 > 1e-5:
            fail("phi=90 mismatch to actual prescan at z={} delta={} mN".format(z, delta90))
        peak = max(ordered, key=lambda row: num(row, "Fx_total_corr_mN"))
        trough = min(ordered, key=lambda row: num(row, "Fx_total_corr_mN"))
        fmax = num(peak, "Fx_total_corr_mN")
        summary_row = {
            "target_force_mN": target["target_force_mN"],
            "reference_z_mm": target["reference_z_mm"],
            "z_sphere_mm": z,
            "alpha_deg": -30.0,
            "prescan_phi90_Fx_mN": num(prescan, "Fx_total_corr_mN"),
            "selection_abs_error_mN": abs(num(prescan, "Fx_total_corr_mN")-target["target_force_mN"]),
            "Fmax_mN": fmax,
            "phi_at_Fmax_deg": num(peak, "phi_deg"),
            "Fmax_target_abs_error_mN": abs(fmax-target["target_force_mN"]),
            "Fmin_mN": num(trough, "Fx_total_corr_mN"),
            "phi_at_Fmin_deg": num(trough, "phi_deg"),
            "positive_phase_count": sum(num(row, "Fx_total_corr_mN") > 0 for row in ordered),
            "negative_phase_count": sum(num(row, "Fx_total_corr_mN") < 0 for row in ordered),
            "closure_max_mN": closure_max,
            "closure_Fx_mN": closure_values["Fx_total_corr_mN"],
            "closure_DeltaFx_mN": closure_values["DeltaFx_ball_mN"],
            "closure_Fy_mN": closure_values["Fy_total_corr_mN"],
            "closure_Fz_mN": closure_values["Fz_total_corr_mN"],
            "phase_points": len(ordered),
            "status": "SUCCESS",
        }
        summary_rows.append(summary_row)

    prescan_out = DATA / "modelB_alpha_minus30_target_prescan.csv"
    selected_out = DATA / "modelB_alpha_minus30_selected_targets.csv"
    full_out = DATA / "modelB_alpha_minus30_target_full360.csv"
    summary_out = DATA / "modelB_alpha_minus30_target_full360_summary.csv"
    xlsx_out = DATA / "modelB_alpha_minus30_selected4_full360.xlsx"
    for dest, source in ((prescan_out, None), (selected_out, args.selected_csv), (full_out, args.full_csv)):
        if dest.exists():
            fail("refusing to overwrite existing result {}".format(dest))
        if source is not None:
            shutil.copyfile(source, dest)
    with prescan_out.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=PRESCAN_COLUMNS)
        writer.writeheader()
        writer.writerows(all_pre)
    with summary_out.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(summary_rows[0]))
        writer.writeheader()
        writer.writerows(summary_rows)

    book = Workbook()
    sheet = book.active
    sheet.title = "full360"
    sheet.append(EXCEL_COLUMNS)
    for z in z_order:
        for row in sorted(groups[z], key=lambda item: num(item, "phi_deg")):
            sheet.append(excel_values(row))
    book.save(xlsx_out)
    book.close()
    validate_workbook(xlsx_out, z_order)
    print("ALPHA_MINUS30_FINAL_VALID rows=404 groups=4 sheets=full360")
    for row in summary_rows:
        print("target={:.9f} mN z={:g} mm Fmax={:.12g} mN @ {:g} deg Fmin={:.12g} mN closure_max={:.4g} mN".format(
            row["target_force_mN"], row["z_sphere_mm"], row["Fmax_mN"], row["phi_at_Fmax_deg"], row["Fmin_mN"], row["closure_max_mN"]))
    print("CSV={}".format(full_out))
    print("XLSX={}".format(xlsx_out))


def main() -> None:
    parser = argparse.ArgumentParser()
    modes = parser.add_mutually_exclusive_group(required=True)
    modes.add_argument("--preflight", action="store_true")
    modes.add_argument("--propose-refinement", action="store_true")
    modes.add_argument("--select", action="store_true")
    modes.add_argument("--finalize", action="store_true")
    parser.add_argument("--coarse-csv", type=Path)
    parser.add_argument("--refine-csv", type=Path)
    parser.add_argument("--proposal-csv", type=Path)
    parser.add_argument("--selected-csv", type=Path)
    parser.add_argument("--full-csv", type=Path)
    parser.add_argument("--run-summary", type=Path)
    parser.add_argument("--candidate-xlsx", type=Path)
    args = parser.parse_args()

    if args.preflight:
        targets = target_levels()
        outputs = [
            DATA / "modelB_alpha_minus30_target_prescan.csv",
            DATA / "modelB_alpha_minus30_selected_targets.csv",
            DATA / "modelB_alpha_minus30_target_full360.csv",
            DATA / "modelB_alpha_minus30_target_full360_summary.csv",
            DATA / "modelB_alpha_minus30_selected4_full360.xlsx",
        ]
        existing = [str(path) for path in outputs if path.exists()]
        if existing:
            fail("alpha=-30 result exists; refusing overwrite: {}".format(existing))
        if len(targets) != 4:
            fail("expected four existing alpha=-20 Fmax targets")
        print("PREFLIGHT_OK targets_mN={}".format([round(row["target_force_mN"], 12) for row in targets]))
        print("COARSE_Z_GRID={}".format(COARSE_ZS))
    elif args.propose_refinement:
        if args.coarse_csv is None or args.proposal_csv is None:
            fail("--propose-refinement needs --coarse-csv and --proposal-csv")
        propose_refinement(args.coarse_csv, args.proposal_csv)
    elif args.select:
        if args.coarse_csv is None or args.refine_csv is None or args.selected_csv is None:
            fail("--select needs --coarse-csv, --refine-csv, and --selected-csv")
        rows = combine_prescans(args.coarse_csv, args.refine_csv)
        select_measured_heights(rows, args.selected_csv)
    else:
        required = [args.coarse_csv, args.refine_csv, args.selected_csv, args.full_csv, args.run_summary]
        if any(path is None for path in required):
            fail("--finalize needs all prescan, selection, and full360 CSV paths")
        finalize(args)


if __name__ == "__main__":
    main()
